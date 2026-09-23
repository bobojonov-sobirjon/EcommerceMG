"""Выгрузка товаров в Excel (.xlsx) или CSV для клиента."""
from __future__ import annotations

import csv
import os
from datetime import datetime
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand

from commerce.models import Product, ProductType


HEADERS = [
    'ID',
    'Артикул',
    'Наименование',
    'Тип товара (код)',
    'Тип товара',
    'Производитель',
    'Цена с НДС',
    'Цена по запросу',
    'В наличии',
    'Slug',
    'SEO заголовок',
    'SEO описание',
    'Описание',
    'Порядок',
    'Кол-во фото',
    'Главное фото (ссылка)',
    'Все фото (ссылки)',
    'Создано',
]

# Колонка «Главное фото» — 16-я (P), в Excel делаем кликабельной.
MAIN_PHOTO_COL = 16


def _absolute_media_url(file_field, base_url: str) -> str:
    if not file_field or not file_field.name:
        return ''
    url = file_field.url
    if url.startswith('http://') or url.startswith('https://'):
        return url
    return f'{base_url.rstrip("/")}{url}'


def _image_urls(product: Product, base_url: str) -> list[str]:
    images = sorted(product.images.all(), key=lambda img: (img.ordering, img.pk))
    urls = []
    for img in images:
        url = _absolute_media_url(img.image, base_url)
        if url:
            urls.append(url)
    return urls


def _product_rows(qs, base_url: str):
    for p in qs:
        seo = getattr(p, 'seo_record', None)
        urls = _image_urls(p, base_url)
        yield [
            p.pk,
            p.artikul,
            p.name,
            p.product_type,
            p.get_product_type_display(),
            p.manufacturer.name if p.manufacturer_id else '',
            '' if p.price_on_request else str(p.price),
            'да' if p.price_on_request else 'нет',
            'да' if p.is_stock else 'нет',
            seo.slug if seo else '',
            seo.seo_title if seo else '',
            seo.seo_description if seo else '',
            p.description or '',
            p.ordering,
            len(urls),
            urls[0] if urls else '',
            '\n'.join(urls),
            p.created_at.strftime('%Y-%m-%d %H:%M') if p.created_at else '',
        ]


class Command(BaseCommand):
    help = 'Выгружает товары в файл .xlsx (Excel) или .csv'

    def add_arguments(self, parser):
        parser.add_argument(
            '--format',
            choices=('xlsx', 'csv'),
            default='xlsx',
            help='Формат файла (по умолчанию xlsx)',
        )
        parser.add_argument(
            '--type',
            choices=[c[0] for c in ProductType.choices],
            default=None,
            help='Фильтр по типу: spare_parts | tires | engines',
        )
        parser.add_argument(
            '--output',
            type=str,
            default='',
            help='Путь к файлу (по умолчанию media/exports/...)',
        )
        parser.add_argument(
            '--base-url',
            type=str,
            default='',
            help='Домен для ссылок на фото (по умолчанию PUBLIC_SITE_URL или https://admin.maksan-group.ru)',
        )

    def handle(self, *args, **options):
        fmt: str = options['format']
        ptype = options['type']
        output = (options['output'] or '').strip()
        base_url = (
            (options['base_url'] or '').strip()
            or os.getenv('PUBLIC_SITE_URL', '').strip()
            or 'https://admin.maksan-group.ru'
        ).rstrip('/')

        qs = (
            Product.objects.select_related('manufacturer', 'seo_record')
            .prefetch_related('images')
            .order_by('product_type', 'manufacturer__name', 'artikul', 'id')
        )
        if ptype:
            qs = qs.filter(product_type=ptype)

        total = qs.count()
        if total == 0:
            self.stdout.write(self.style.WARNING('Товаров не найдено.'))
            return

        stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        type_part = ptype or 'all'
        if not output:
            export_dir = Path(settings.MEDIA_ROOT) / 'exports'
            export_dir.mkdir(parents=True, exist_ok=True)
            output = str(export_dir / f'products_{type_part}_{stamp}.{fmt}')

        path = Path(output)
        path.parent.mkdir(parents=True, exist_ok=True)

        rows = list(_product_rows(qs, base_url))

        if fmt == 'xlsx':
            self._write_xlsx(path, rows)
        else:
            self._write_csv(path, rows)

        self.stdout.write(self.style.SUCCESS(f'Готово: {total} товаров -> {path.resolve()}'))

    def _write_csv(self, path: Path, rows: list) -> None:
        with path.open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f, delimiter=';')
            writer.writerow(HEADERS)
            writer.writerows(rows)

    def _write_xlsx(self, path: Path, rows: list) -> None:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font
        except ImportError:
            self.stdout.write(
                self.style.WARNING('openpyxl не установлен — сохраняю CSV. pip install openpyxl'),
            )
            csv_path = path.with_suffix('.csv')
            self._write_csv(csv_path, rows)
            self.stdout.write(self.style.SUCCESS(f'CSV: {csv_path.resolve()}'))
            raise SystemExit(0)

        wb = Workbook()
        ws = wb.active
        ws.title = 'Товары'
        ws.append(HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        link_font = Font(color='0563C1', underline='single')
        wrap = Alignment(wrap_text=True, vertical='top')
        for row in rows:
            ws.append(row)
            excel_row = ws.max_row
            photo_cell = ws.cell(row=excel_row, column=MAIN_PHOTO_COL)
            if photo_cell.value:
                photo_cell.hyperlink = str(photo_cell.value)
                photo_cell.font = link_font
            all_photos = ws.cell(row=excel_row, column=MAIN_PHOTO_COL + 1)
            all_photos.alignment = wrap

        widths = {
            'A': 8,
            'B': 14,
            'C': 48,
            'D': 14,
            'E': 22,
            'F': 18,
            'G': 14,
            'H': 14,
            'I': 12,
            'J': 28,
            'K': 28,
            'L': 36,
            'M': 40,
            'N': 10,
            'O': 12,
            'P': 55,
            'Q': 55,
            'R': 18,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        wb.save(path)
